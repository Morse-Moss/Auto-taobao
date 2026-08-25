#!/usr/bin/env python3
"""Verify that a SYCM CSV and XLSX are the same proven seven-day export."""

from __future__ import annotations

import csv
import hashlib
import json
import sys
from datetime import date
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook


HEADERS = ["排名", "搜索词", "搜索人气", "点击率", "支付转化率"]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        rows = list(csv.reader(source))
    if not rows or rows[0] != HEADERS:
        raise ValueError(f"CSV headers must be exactly: {','.join(HEADERS)}")
    return rows


def read_workbook(path: Path) -> tuple[list[list[str]], dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if workbook.sheetnames != ["搜索排行", "验证信息"]:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    data_rows = [[cell_text(cell.value) for cell in row] for row in workbook["搜索排行"].iter_rows()]
    if not data_rows or data_rows[0] != HEADERS:
        raise ValueError(f"XLSX headers must be exactly: {','.join(HEADERS)}")
    metadata: dict[str, object] = {}
    for key_cell, value_cell in workbook["验证信息"].iter_rows(min_row=2, max_col=2):
        key = cell_text(key_cell.value)
        if key:
            metadata[key] = value_cell.value
    return data_rows, metadata


def verify(csv_path: Path, xlsx_path: Path, expected_end_date: str) -> dict[str, object]:
    csv_rows = read_csv(csv_path)
    workbook_rows, metadata = read_workbook(xlsx_path)
    if csv_rows != workbook_rows:
        raise ValueError("CSV and XLSX data differ")

    start_date = cell_text(metadata.get("startDate"))
    end_date = cell_text(metadata.get("endDate"))
    period = cell_text(metadata.get("period"))
    day_count = int(metadata.get("dayCount", 0))
    start = date.fromisoformat(start_date)
    end = date.fromisoformat(end_date)
    if period != "7天" or day_count != 7 or (end - start).days != 6 or end_date != expected_end_date:
        raise ValueError("Export pair is not a verified 7-day reporting window ending on the collection date")

    source_query = parse_qs(urlparse(cell_text(metadata.get("sourceUrl"))).query)
    expected_range = f"{start_date}|{end_date}"
    if source_query.get("dateType") != ["recent7"] or source_query.get("dateRange") != [expected_range]:
        raise ValueError("Workbook source URL does not prove the same seven-day reporting window")

    row_count = len(csv_rows) - 1
    if int(metadata.get("rowCount", 0)) != row_count:
        raise ValueError("Workbook validation row count does not match the export")
    for name in ("uniqueRanks", "contiguousRanks", "uniqueTerms", "nonEmptyMetrics"):
        if metadata.get(name) is not True:
            raise ValueError(f"Workbook validation did not prove {name}")

    return {
        "period": period,
        "startDate": start_date,
        "endDate": end_date,
        "dayCount": day_count,
        "dateRange": cell_text(metadata.get("dateRange")),
        "rowCount": row_count,
        "rankRange": cell_text(metadata.get("rankRange")),
        "pageSizes": cell_text(metadata.get("pageSizes")),
        "csvSha256": sha256(csv_path),
        "xlsxSha256": sha256(xlsx_path),
    }


def main() -> int:
    if len(sys.argv) != 4:
        print("Usage: verify-export-pair.py SOURCE.csv SOURCE.xlsx EXPECTED_END_DATE", file=sys.stderr)
        return 2
    try:
        result = verify(Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3])
    except Exception as exc:
        print(json.dumps({"status": "error", "message": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps({"status": "success", **result}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
