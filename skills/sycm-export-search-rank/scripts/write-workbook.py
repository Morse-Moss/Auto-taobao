#!/usr/bin/env python3
"""Write and validate the Feishu-ready SYCM workbook."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def write_workbook(payload_path: Path, output_path: Path) -> dict:
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    metadata = payload["metadata"]
    rows = payload["rows"]
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "搜索排行"
    headers = ["排名", "搜索词", "搜索人气", "点击率", "支付转化率"]
    data_sheet.append(headers)
    for row in rows:
        data_sheet.append([
            row["rank"],
            row["term"],
            row["searchPopularity"],
            row["clickRate"],
            row["payConversionRate"],
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="000000")
    for cell in data_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in data_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column == 2)

    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = f"A1:E{len(rows) + 1}"
    data_sheet.column_dimensions["A"].width = 10
    data_sheet.column_dimensions["B"].width = 32
    for column in ("C", "D", "E"):
        data_sheet.column_dimensions[column].width = 18
    data_sheet.row_dimensions[1].height = 22
    if rows:
        table = Table(displayName="SearchRankData", ref=f"A1:E{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        data_sheet.add_table(table)

    info_sheet = workbook.create_sheet("验证信息")
    info_sheet.append(["字段", "值"])
    for cell in info_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for key, value in metadata.items():
        if isinstance(value, list):
            value = "/".join(str(item) for item in value)
        info_sheet.append([key, value])
    for row in info_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    info_sheet.freeze_panes = "A2"
    info_sheet.column_dimensions["A"].width = 24
    info_sheet.column_dimensions["B"].width = 72
    info_sheet.row_dimensions[1].height = 22

    workbook.save(output_path)

    check = load_workbook(output_path, data_only=False)
    errors = []
    for sheet in check.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    if check.sheetnames != ["搜索排行", "验证信息"]:
        raise ValueError(f"Unexpected workbook sheets: {check.sheetnames}")
    if check["搜索排行"].max_row != len(rows) + 1:
        raise ValueError("Workbook row count does not match payload")
    if errors:
        raise ValueError(f"Formula/error cells found: {errors[:5]}")
    return {"status": "success", "formulaErrors": 0, "rows": len(rows), "sheets": check.sheetnames}


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: write-workbook.py PAYLOAD.json OUTPUT.xlsx", file=sys.stderr)
        return 2
    try:
        result = write_workbook(Path(sys.argv[1]), Path(sys.argv[2]))
    except Exception as exc:  # Report a concise machine-readable failure to the caller.
        print(json.dumps({"status": "error", "message": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
