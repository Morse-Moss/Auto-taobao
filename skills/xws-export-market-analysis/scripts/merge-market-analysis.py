#!/usr/bin/env python3

import argparse
import csv
import io
import json
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image


def read_csv(path):
    last_error = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with open(path, "r", encoding=encoding, newline="") as handle:
                return list(csv.reader(handle))
        except UnicodeDecodeError as error:
            last_error = error
    raise last_error


def image_map(path):
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    values = list(sheet.iter_rows(values_only=True))
    result = {}
    for image in sheet._images:
        anchor = image.anchor
        row_index = getattr(getattr(anchor, "_from", None), "row", None)
        if row_index is None or row_index >= len(values):
            continue
        row = values[row_index]
        link = str(row[3] or "").strip() if len(row) > 3 else ""
        if not link or link in result:
            continue
        data = image._data()
        with Image.open(io.BytesIO(data)) as source:
            source.verify()
        result[link] = (data, image.width, image.height)
    workbook.close()
    return result


def write_xlsx(csv_path, source_paths, output_path, require_images):
    table = read_csv(csv_path)
    if not table:
        raise ValueError("merged CSV is empty")
    headers, rows = table[0], table[1:]
    images = {}
    for source in source_paths:
        images.update({key: value for key, value in image_map(source).items() if key not in images})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "市场分析"
    sheet.append(headers)
    missing = []
    temporary_images = []
    try:
        for index, row in enumerate(rows, 2):
            sheet.append(row)
            link = str(row[3] or "").strip()
            image = images.get(link)
            if image is None:
                missing.append(link)
                continue
            image_data, width, height = image
            temporary = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            temporary.write(image_data)
            temporary.close()
            temporary_images.append(temporary.name)
            worksheet_image = WorksheetImage(temporary.name)
            worksheet_image.width = width
            worksheet_image.height = height
            sheet.add_image(worksheet_image, f"B{index}")
        # 图片缺口判据（2026-09-15 与 validate-output.py 改成同一口径，两处必须一起改）：
        #   历史缺陷的样子是「整本没有图」→ 仍然失败；要求「一张不缺」会把**单张图抓取失败**
        #   变成整条周更流程的硬失败（2026-09-15 实测：166 行 / 165 图，白等 40 分钟后
        #   在采集端被拒；若只在采集端放宽，就会把失败推到这里）。
        #   所以：0 张 → 失败；缺口超过 2%（至少容 1 张）→ 失败；其余接受，
        #   缺口由返回值里的 missingImages 如实上报，不静默吞掉。
        tolerance = max(1, (len(rows) + 49) // 50)
        if require_images and (len(rows) - len(missing) == 0 or len(missing) > tolerance):
            raise ValueError(f"missing images for {len(missing)} merged rows (tolerance {tolerance})")
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return {"path": str(Path(output_path).resolve()), "rows": len(rows), "images": len(rows) - len(missing), "missingImages": len(missing)}
    finally:
        workbook.close()
        for temporary in temporary_images:
            try:
                os.unlink(temporary)
            except FileNotFoundError:
                pass


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv", required=True)
    parser.add_argument("--source-xlsx", action="append", default=[])
    parser.add_argument("--output-xlsx", required=True)
    parser.add_argument("--require-images", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()
    try:
        result = write_xlsx(args.csv, args.source_xlsx, args.output_xlsx, args.require_images)
        print(json.dumps({"ok": True, **result}, ensure_ascii=False))
        return 0
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
