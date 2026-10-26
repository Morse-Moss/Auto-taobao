#!/usr/bin/env python3

import argparse
import csv
import io
import json
import zipfile
from datetime import date, datetime

import openpyxl


# 推广 CSV 允许的两种真实宽度（按店不同，见下方注释）。形状校验只认这两个值。
PROMOTION_COLUMN_WIDTHS = (71, 73)


def as_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or '').strip()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--shop-xlsx', required=True)
    parser.add_argument('--promotion-zip', required=True)
    parser.add_argument('--date', required=True)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.shop_xlsx, data_only=True, read_only=True)
    if workbook.sheetnames != ['data']:
        raise ValueError(f'expected only data sheet, got {workbook.sheetnames}')
    rows = list(workbook['data'].iter_rows(values_only=True))
    if not rows or len(rows[0]) != 119:
        raise ValueError(f'expected 119 shop columns, got {len(rows[0]) if rows else 0}')
    matches = [row for row in rows[1:] if as_date(row[0]) == args.date]
    if len(matches) != 1:
        raise ValueError(f'expected one shop row for {args.date}, got {len(matches)}')

    with zipfile.ZipFile(args.promotion_zip) as archive:
        names = [name for name in archive.namelist() if not name.endswith('/')]
        if len(names) != 1 or not names[0].lower().endswith('.csv'):
            raise ValueError(f'expected one CSV in promotion ZIP, got {names}')
        text = archive.read(names[0]).decode('gb18030').lstrip('\ufeff')
    promotion_rows = [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]
    # 推广 CSV 有两种**真实**宽度，按店不同（2026-10-26 实测：同一天 3 份 71 列、保拉天猫 73 列，
    # 差别是有没有开通「补贴」那套指标）。所以这里只做**形状**校验，不钉单值：
    # 钉死 71 会把 73 列的店挡在门外（就是当天 push 失败的那一句）；钉死 73 会把 71 列的店挡在门外。
    # 列的含义逐列核对在 daily-report-core.mjs 的 buildPromotionFields 里做（那里有真实表头可比）。
    width = len(promotion_rows[0]) if promotion_rows else 0
    if width not in PROMOTION_COLUMN_WIDTHS:
        raise ValueError(f'expected one of {PROMOTION_COLUMN_WIDTHS} promotion columns, got {width}')
    if any(len(row) != width for row in promotion_rows):
        raise ValueError('promotion CSV contains a row with an unexpected column count')

    print(json.dumps({
        'shop': {
            'headers': list(rows[0]),
            'values': list(matches[0]),
            'workbookRows': len(rows),
            # 整列日期都带出来：收据要能自证「这份工作簿覆盖哪几天、目标日出现几次」。
            # 光记文件哈希只能证明「是这个文件」，证明不了「是这个文件里的一天」。
            'dates': [as_date(row[0]) for row in rows[1:]],
        },
        'promotion': {
            'headers': promotion_rows[0],
            'rows': promotion_rows[1:],
            'csvName': names[0],
            'dates': [row[0] for row in promotion_rows[1:]],
        },
    }, ensure_ascii=False, default=as_date))


if __name__ == '__main__':
    main()
