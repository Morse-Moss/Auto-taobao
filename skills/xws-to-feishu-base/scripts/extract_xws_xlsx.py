import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def serialize_stdout(value: dict) -> str:
    return json.dumps(value, ensure_ascii=True)


def extract_workbook(workbook_path: Path, output_dir: Path) -> dict:
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook.active
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    if any(not isinstance(header, str) or not header for header in headers):
        raise ValueError("XLSX header row contains an empty or invalid field name")

    rows = [
        {header: sheet.cell(row, column).value for column, header in enumerate(headers, 1)}
        for row in range(2, sheet.max_row + 1)
    ]
    output_dir.mkdir(parents=True, exist_ok=True)
    images = []
    for index, image in enumerate(sheet._images, 1):
        anchor = getattr(image.anchor, "_from", None)
        if anchor is None:
            raise ValueError(f"Image {index} has no worksheet anchor")
        row = anchor.row + 1
        extension = (image.format or "png").lower()
        path = output_dir / f"row-{row:06d}-{index:03d}.{extension}"
        path.write_bytes(image._data())
        images.append({"row": row, "column": anchor.col + 1, "path": str(path.resolve())})

    return {
        "workbook": str(workbook_path.resolve()),
        "sheet": sheet.title,
        "headers": headers,
        "rows": rows,
        "images": images,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--manifest", type=Path)
    args = parser.parse_args()
    result = extract_workbook(args.xlsx, args.output_dir)
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.manifest:
        args.manifest.parent.mkdir(parents=True, exist_ok=True)
        args.manifest.write_text(payload, encoding="utf-8")
    print(serialize_stdout(result))


if __name__ == "__main__":
    main()
