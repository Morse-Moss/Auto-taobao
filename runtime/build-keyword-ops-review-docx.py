from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run


PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

ASCII_FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft YaHei"  # Named override for readable Chinese glyphs.
MONO_FONT = "Consolas"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B3A"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "B7C3D0"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"


def set_run_font(run, *, name=ASCII_FONT, east_asia=EAST_ASIA_FONT, size=None,
                 bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border(paragraph, *, side="left", color=BLUE, size=18, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    tag = qn(f"w:{side}")
    edge = p_bdr.find(tag)
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        p_bdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in CELL_MARGIN_DXA.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), BORDER)


def set_table_geometry(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def distribute_widths(rows):
    column_count = len(rows[0])
    raw = []
    for col in range(column_count):
        lengths = [len(re.sub(r"[`*]", "", row[col])) for row in rows]
        raw.append(max(8, min(36, max(lengths, default=8))))
    if column_count >= 4:
        raw[0] = min(raw[0], 14)
    total = sum(raw)
    widths = [max(1050, round(CONTENT_WIDTH_DXA * value / total)) for value in raw]
    while sum(widths) > CONTENT_WIDTH_DXA:
        idx = max(range(column_count), key=lambda i: widths[i])
        widths[idx] -= 1
    while sum(widths) < CONTENT_WIDTH_DXA:
        idx = max(range(column_count), key=lambda i: raw[i])
        widths[idx] += 1
    return widths


def add_hyperlink(paragraph, text, url, *, size=11):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)
    run = Run(run_element, paragraph)
    run.text = text
    set_run_font(run, size=size, color=BLUE)
    run.font.underline = True


def add_inline(paragraph, text, *, size=11, color=INK, bold=False, italic=False):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|https?://[^\s]+)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name=MONO_FONT, size=max(9, size - 0.5), color=DARK_BLUE)
            rpr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_GRAY)
            rpr.append(shd)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif part.startswith(("http://", "https://")):
            add_hyperlink(paragraph, part, part, size=size)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color, bold=bold, italic=italic)


def append_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    run.extend([begin, instr, separate, text, end])
    paragraph._p.append(run)


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = ASCII_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = ASCII_FONT
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part._element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.2


def add_metadata_table(doc, metadata):
    table = doc.add_table(rows=len(metadata), cols=2)
    widths = [1500, CONTENT_WIDTH_DXA - 1500]
    set_table_geometry(table, widths)
    for row, (label, value) in zip(table.rows, metadata):
        for cell in row.cells:
            set_cell_shading(cell, WHITE)
        label_p = row.cells[0].paragraphs[0]
        label_p.paragraph_format.space_after = Pt(2)
        add_inline(label_p, label, size=10, color=MUTED, bold=True)
        value_p = row.cells[1].paragraphs[0]
        value_p.paragraph_format.space_after = Pt(2)
        add_inline(value_p, value, size=10, color=INK)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    tbl_pr.append(borders)


def add_title_block(doc, title, subtitle, metadata):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(3)
    add_inline(kicker, "运营审核材料 · V1", size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    add_inline(p, title, size=24, color=INK, bold=True)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(12)
    sub.paragraph_format.keep_with_next = True
    add_inline(sub, subtitle, size=12.5, color=MUTED)

    add_metadata_table(doc, metadata)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(4)
    rule.paragraph_format.space_after = Pt(10)
    set_paragraph_border(rule, side="bottom", color=BLUE, size=12, space=2)


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_shading(p, CALLOUT)
    set_paragraph_border(p, side="left", color=BLUE, size=20, space=8)
    add_inline(p, text, size=10.5, color=INK, bold=True)


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_together = True
    set_paragraph_shading(p, LIGHT_GRAY)
    set_paragraph_border(p, side="left", color=DARK_BLUE, size=12, space=7)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name=MONO_FONT, east_asia=EAST_ASIA_FONT, size=9.2, color=DARK_BLUE)


def clean_table_cell(text):
    text = text.strip()
    text = re.sub(r"^`|`$", "", text)
    return text.replace(r"\|", "|")


def add_markdown_table(doc, rows):
    rows = [[clean_table_cell(cell) for cell in row] for row in rows]
    column_count = len(rows[0])
    rows = [row + [""] * (column_count - len(row)) for row in rows]
    widths = distribute_widths(rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.autofit = False
    set_table_geometry(table, widths)
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for r_idx, (word_row, values) in enumerate(zip(table.rows, rows)):
        for c_idx, (cell, value) in enumerate(zip(word_row.cells, values)):
            set_cell_shading(cell, LIGHT_BLUE if r_idx == 0 else WHITE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or re.fullmatch(r"[0-9]+", value)) else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, size=9.2 if column_count <= 4 else 8.7, color=INK, bold=(r_idx == 0))
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(1)


def is_table_separator(line):
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def setup_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, "关键词分析 V1  |  双表运营审核", size=8.5, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, "第 ", size=8.5, color=MUTED)
    append_field(p, "PAGE")
    add_inline(p, " 页", size=8.5, color=MUTED)


def render_markdown(doc, lines, start_index, bullet_num_id):
    idx = start_index
    decimal_num_id = None
    in_numbered_list = False
    while idx < len(lines):
        line = lines[idx].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            in_numbered_list = False
            idx += 1
            continue

        if stripped.startswith("```"):
            in_numbered_list = False
            idx += 1
            code_lines = []
            while idx < len(lines) and not lines[idx].strip().startswith("```"):
                code_lines.append(lines[idx].rstrip())
                idx += 1
            add_code_block(doc, code_lines)
            idx += 1
            continue

        if stripped.startswith("|") and idx + 1 < len(lines) and is_table_separator(lines[idx + 1]):
            in_numbered_list = False
            rows = [split_table_row(stripped)]
            idx += 2
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                rows.append(split_table_row(lines[idx]))
                idx += 1
            add_markdown_table(doc, rows)
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            in_numbered_list = False
            level = min(len(heading.group(1)) - 1, 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, heading.group(2), size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
            idx += 1
            continue

        if stripped.startswith(">"):
            in_numbered_list = False
            add_callout(doc, stripped[1:].strip())
            idx += 1
            continue

        if stripped.startswith("- "):
            in_numbered_list = False
            p = doc.add_paragraph()
            apply_numbering(p, bullet_num_id)
            add_inline(p, stripped[2:], size=11, color=INK)
            idx += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            if not in_numbered_list:
                decimal_num_id = add_numbering_definition(doc, "decimal")
                in_numbered_list = True
            p = doc.add_paragraph()
            apply_numbering(p, decimal_num_id)
            add_inline(p, numbered.group(1), size=11, color=INK)
            idx += 1
            continue

        in_numbered_list = False
        p = doc.add_paragraph()
        add_inline(p, stripped, size=11, color=INK)
        idx += 1


def build(source, output):
    lines = source.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("# "):
        raise ValueError("Source must start with a level-one title")

    title = lines[0][2:].strip()
    subtitle = lines[2].split("：", 1)[1].strip()
    metadata = []
    for line in lines[3:6]:
        label, value = line.split("：", 1)
        metadata.append((label, value))

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    setup_header_footer(section)

    doc.core_properties.title = title
    doc.core_properties.subject = "浴缸关键词双表运营审核合同 V1"
    doc.core_properties.author = "sycm-automation"
    doc.core_properties.keywords = "淘宝, 生意参谋, 浴缸, 关键词, 飞书, 运营审核"

    bullet_num_id = add_numbering_definition(doc, "bullet")
    add_title_block(doc, title, subtitle, metadata)
    render_markdown(doc, lines, 7, bullet_num_id)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()
