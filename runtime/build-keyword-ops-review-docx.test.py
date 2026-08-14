from pathlib import Path
import importlib.util
import unittest

from docx import Document


module_path = Path(__file__).with_name("build-keyword-ops-review-docx.py")
spec = importlib.util.spec_from_file_location("build_keyword_ops_review_docx", module_path)
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


class HyperlinkRenderingTest(unittest.TestCase):
    def test_add_inline_turns_http_url_into_a_real_hyperlink(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        url = "https://example.com/base?table=tbl123&view=vew456"

        builder.add_inline(paragraph, f"打开 {url}")

        links = [
            relationship.target_ref
            for relationship in doc.part.rels.values()
            if relationship.reltype.endswith("/hyperlink")
        ]
        self.assertIn(url, links)

    def test_add_inline_writes_visible_hyperlink_text(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        url = "https://example.com/base?table=tbl123&view=vew456"

        builder.add_inline(paragraph, url)

        hyperlink = paragraph._p.xpath("./w:hyperlink")[0]
        visible_text = [node.text for node in hyperlink.xpath("./w:r/w:t")]
        self.assertEqual([url], visible_text)


if __name__ == "__main__":
    unittest.main()
