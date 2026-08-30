from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_text(document, text, *, bold=False, size=10.5):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    return paragraph


def qn(value):
    from docx.oxml.ns import qn as qualified_name
    return qualified_name(value)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: generate-weekly-analysis-doc.py MODEL_JSON OUTPUT_DOCX")
    model = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    output = Path(sys.argv[2])
    document = Document()
    document.core_properties.title = model["title"]
    document.core_properties.subject = "本地分析与飞书最终值发布"

    title = document.add_heading(model["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"状态：{model['status']}  |  注册表版本：{model['registryVersion']}  |  注册表 SHA-256：{model['registryDigest']}")
    document.add_paragraph(f"Artifact SHA-256：{model['artifactDigest']}  |  Publish plan SHA-256：{model['planDigest']}")
    source = model['sourceEvidence']
    digests = model['evidenceDigests']
    document.add_paragraph(f"SYCM CSV SHA-256：{source.get('csvSha256', '')}  |  XLSX SHA-256：{source.get('xlsxSha256', '')}  |  输入快照 digest：{source.get('inputSnapshotDigest', '')}")
    document.add_paragraph(f"Provider digest：{digests['provider']}  |  Prompt digest：{digests['prompt']}  |  灰豚 digest：{digests['huitun']}  |  历史快照 digest：{digests['history']}  |  编号库快照 digest：{digests['library']}")
    document.add_heading("运行摘要", 1)
    provider = model['providerSummary']
    huitun = model['huitunSummary']
    publish = model['publishSummary']
    document.add_paragraph(f"Provider：{', '.join(provider['providers']) or '未指定'}；已验证结果：{provider['validated']} / {provider['total']}。")
    document.add_paragraph(f"灰豚证据条目：{huitun['items']}；采集时间：{huitun['collectedAt'] or '未提供'}。")
    document.add_paragraph(f"发布计划：当前周表更新 {publish['currentUpdates']} 条；历史表新增 {publish['historyCreates']} 条、补齐 {publish['historyUpdates']} 条；编号库新增 {publish['libraryCreates']} 条。")
    document.add_heading("运行边界", 1)
    document.add_paragraph(model["evidenceBoundary"]["feishu"])
    document.add_paragraph(model["missingValuePolicy"])

    document.add_heading("字段 Owner 矩阵", 1)
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["字段", "Owner", "输入", "输出", "缺失语义"]):
        cell.text = value
    for field in model["fields"]:
        row = table.add_row().cells
        row[0].text = field["name"]
        row[1].text = field["owner"]
        row[2].text = "、".join(field["input"])
        row[3].text = field["output"]
        row[4].text = field["missing"]

    document.add_heading("Prompt", 1)
    document.add_paragraph("内容热度")
    document.add_paragraph(model["prompts"]["contentHeat"])
    document.add_paragraph("对应产品方向")
    document.add_paragraph(model["prompts"]["productDirection"])

    document.add_heading("外部证据边界", 1)
    document.add_paragraph(model["evidenceBoundary"]["sycm"])
    document.add_paragraph(model["evidenceBoundary"]["huitun"])

    document.add_heading("发布原则", 1)
    document.add_paragraph("本地阶段生成并校验 analysis artifact、provider 结果、灰豚证据和 publish plan；post 阶段只在目标 Base、目标表、记录身份和 plan digest 全部一致时批量回传，并在回传后回读验证。")
    document.save(output)


if __name__ == "__main__":
    main()
