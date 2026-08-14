from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PATH = Path(r"C:\Users\Administrator\Desktop\浴缸关键词经营决策三字段说明与阶段成果汇报_V1.0_20260813.docx")
BLUE = "2E74B5"
NAVY = "17324D"
TEXT = "202124"
MUTED = "5F6368"


def set_font(run, size=10.3, bold=False, color=TEXT, code=False):
    family = "Consolas" if code else "Calibri"
    run.font.name = family
    props = run._element.get_or_add_rPr()
    props.rFonts.set(qn("w:ascii"), family)
    props.rFonts.set(qn("w:hAnsi"), family)
    props.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, size=10.3, bold=False, color=TEXT, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run(text), size, bold, color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(text), 16 if level == 1 else 12.5, True, BLUE if level == 1 else NAVY)
    return p


def code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.18)
    p.paragraph_format.right_indent = Inches(.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    ppr = p._p.get_or_add_pPr()
    shd = ppr.makeelement(qn("w:shd"), {qn("w:fill"): "F2F4F7"})
    ppr.append(shd)
    for i, line in enumerate(text.strip().splitlines()):
        if i:
            p.add_run().add_break()
        set_font(p.add_run(line), 8.8, color="263238", code=True)
    return p


def append():
    doc = Document(PATH)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "附录：飞书公式与 AI 提示词")
    add_text(doc, "公式已经填入当前飞书表。下面是字段名展示版，用来说明逻辑和留档。当前表中实际公式已绑定对应字段；复制到新表时，需要重新绑定字段，不能直接原样粘贴。", 10.5, color=MUTED, after=9)

    add_heading(doc, "1. 是否重点词公式", 2)
    add_text(doc, "作用：判断这个词是否是长期、稳定的产品需求。", 10.5, bold=True)
    code_block(doc, '''IF(搜索词为空, "",
  IF(已有有效批次数为空 或 少于8, "观察中",
    IF(关键词分类="品牌词"
       或 细分标签包含"痛点/清洁、痛点/漏水、痛点/排水、痛点/维修", "否",
      IF(近8批出现次数>=6, "是", "否")
    )
  )
)''')
    add_text(doc, "输出只有四种：留空、观察中、是、否。", 10, color=MUTED)

    add_heading(doc, "2. 优先级公式", 2)
    add_text(doc, "作用：判断现在是否需要抢时间跟进。它不读取“是否重点词”的结果。", 10.5, bold=True)
    code_block(doc, '''IF(搜索词为空, "",
  IF(已有有效批次数为空 或 等于0, "观察中",
    IF(品牌词 或 服务型痛点, "C-常规跟踪",
      IF(已有有效批次数>=2 且 近8批出现次数=1 且 排名<=100
         且 用户意图为购买型/对比型 且 搜索热度="高"
         且 交易热度为"中"或"高", "A-立即跟进",
        IF(用户意图为购买型/对比型 且 搜索热度="高"
           且 交易热度为"中"或"高", "B-持续观察",
          "C-常规跟踪"
        )
      )
    )
  )
)''')
    add_text(doc, "输出：A-立即跟进、B-持续观察、C-常规跟踪、观察中。", 10, color=MUTED)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "3. 对应产品方向 AI 提示词", 2)
    add_text(doc, "填写位置：飞书“对应产品方向”AI 字段。输入字段为原始关键词、标准归并词、关键词分类、细分标签。", 10.5, color=MUTED)
    code_block(doc, '''你是浴缸产品方向分析助手。

原始关键词：{{原始关键词}}
标准归并词：{{标准归并词}}
关键词分类：{{关键词分类}}
细分标签：{{细分标签}}

目标：给运营一个可参考的产品方向名称。只输出方向名称，不解释；没有明确方向时留空。

规则：
1. 原始关键词是最终事实依据；标准归并词、关键词分类、细分标签只帮助理解。字段冲突时，以原始关键词为准。
2. 只保留原始关键词明确写出的产品属性。禁止根据常识补充原词没有写出的属性，禁止推断。
3. 保留会改变产品方案的明确属性；相同属性组合统一命名，不同属性不能合并。
4. 命名顺序：定位 + 人群/场景 + 尺寸 + 材质 + 风格 + 形状/款式 + 安装方式 + 功能/需求 + 浴缸。只保留有证据的部分。
5. 品牌、店铺、交易导航、地域、颜色，不进入产品方向。
6. 明确痛点可以转成产品需求，例如“浴缸太滑”输出“防滑浴缸”。清洁、维修、漏水处理等服务问题留空。
7. 老人、户外、小户型等场景，不得擅自补充防滑、耐候、深泡等属性。高端和普通、深泡浴缸和通用浴缸，是不同方向。
8. 只有通用品类“浴缸”且没有其他产品属性时留空。

示例：
家用浴缸 -> 家用浴缸
高端人造石浴缸 -> 高端人造石浴缸
日式深泡小浴缸 -> 小型日式深泡浴缸
科勒浴缸官方旗舰店 -> 留空
浴缸漏水维修 -> 留空

只输出一个方向名称或留空。''')
    add_text(doc, "使用提醒：先抽样运行，重点检查高端、深泡、材质、品牌词和服务词；AI 结果可为空，不要为了填满字段而补造产品属性。", 10.5, bold=True, color=NAVY)
    doc.save(PATH)
    print(PATH)


if __name__ == "__main__":
    append()
