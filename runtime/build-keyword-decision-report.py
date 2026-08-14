from __future__ import annotations

import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(r"D:\Retire\sycm-automation")
OUTPUT = Path(r"C:\Users\Administrator\Desktop\浴缸关键词经营决策三字段说明与阶段成果汇报_V1.0_20260813.docx")
SCREENSHOT = Path(r"C:\Users\ADMINI~1\AppData\Local\Temp\codex-clipboard-8cdae04f-d53f-4af5-bead-6e6a76aa8003.png")
TABLE_HELPER_DIR = Path(r"D:\codex\plugins\cache\openai-primary-runtime\documents\26.812.11052\skills\documents\scripts")
sys.path.insert(0, str(TABLE_HELPER_DIR))
from table_geometry import apply_table_geometry  # noqa: E402


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17324D"
INK = "202124"
GRAY = "5F6368"
MUTED = "7A828A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
PALE_GREEN = "EAF4EA"
PALE_GOLD = "FFF4D6"
PALE_RED = "FCE8E6"
WHITE = "FFFFFF"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 90, "bottom": 90, "start": 120, "end": 120}


def set_run_font(run, *, ascii_font="Calibri", east_asia="Microsoft YaHei", size=None,
                 bold=None, italic=None, color=None):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, *, side="left", color=BLUE, size=12, space=6):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:cantSplit"))
    if marker is None:
        marker = OxmlElement("w:cantSplit")
        tr_pr.append(marker)


def add_page_field(paragraph, field_name):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field_name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run_pr.extend([fonts, color, underline])
    run.append(run_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_numbering_definition(doc, *, ordered):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1." if ordered else "•")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.extend([start, fmt, text, justification])

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, indent, spacing])
    level.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.extend([level, number])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        left = p.add_run("关键词经营决策字段说明")
        set_run_font(left, size=9, color=MUTED, bold=True)
        right = p.add_run("  |  V1.0  |  2026-08-13")
        set_run_font(right, size=9, color=MUTED)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)
        prefix = fp.add_run("第 ")
        set_run_font(prefix, size=9, color=MUTED)
        add_page_field(fp, "PAGE")
        middle = fp.add_run(" 页 / 共 ")
        set_run_font(middle, size=9, color=MUTED)
        add_page_field(fp, "NUMPAGES")
        suffix = fp.add_run(" 页")
        set_run_font(suffix, size=9, color=MUTED)

    props = doc.core_properties
    props.title = "浴缸关键词经营决策三字段说明与阶段成果汇报"
    props.subject = "是否重点词、优先级、对应产品方向及历史辅助字段说明"
    props.author = "关键词分析项目组"
    props.keywords = "飞书多维表格, 浴缸关键词, 重点词, 优先级, 产品方向"
    props.comments = "依据 2026-08-13 已验证字段结构和公式资产生成"


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("阶段成果汇报 / 运营讲解稿")
    set_run_font(r, size=9.5, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("浴缸关键词经营决策三字段说明")
    set_run_font(r, size=25, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("是否重点词 / 优先级 / 对应产品方向")
    set_run_font(r, size=13.5, color=GRAY)

    metadata = [
        ("适用表", "关键词分析 V1（修正版）"),
        ("版本", "V1.0"),
        ("日期", "2026-08-13"),
        ("口径", "运营口径优先；只使用真实采集字段和已确认的 AI 分析字段，不推算、不补造"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}：")
        set_run_font(r, size=10.5, bold=True, color=INK)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("飞书表：")
    set_run_font(r, size=10.5, bold=True, color=INK)
    add_hyperlink(
        p,
        "关键词分析 V1（修正版）",
        "https://rcndesfqro3x.feishu.cn/base/N21Abkg0HakO6AsbCaDckvcwnVd?table=tblN1uT1LpzyqqWx&view=vewwUCNr4y",
    )
    set_paragraph_border(p, side="bottom", color="CBD3DC", size=8, space=8)


def add_callout(doc, label, text, *, fill=CALLOUT, border=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_shading(p, fill)
    set_paragraph_border(p, side="left", color=border, size=18, space=8)
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_label_paragraph(doc, label, text, *, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(f"{label}：")
    set_run_font(r, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, color=INK)
    return p


def add_bullets(doc, items, num_id):
    for item in items:
        p = doc.add_paragraph()
        apply_numbering(p, num_id)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, size=10.8, color=INK)


def add_numbered_steps(doc, items):
    num_id = add_numbering_definition(doc, ordered=True)
    for title, detail in items:
        p = doc.add_paragraph()
        apply_numbering(p, num_id)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(title)
        set_run_font(r, bold=True, color=DARK_BLUE)
        r = p.add_run(f"  {detail}")
        set_run_font(r, color=INK)


def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.0
    set_paragraph_shading(p, LIGHT_GRAY)
    set_paragraph_border(p, side="left", color="96A7B8", size=12, space=6)
    lines = text.strip("\n").split("\n")
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, ascii_font="Consolas", east_asia="Microsoft YaHei", size=9.2, color="263238")
    return p


def add_table(doc, headers, rows, widths, *, header_fill=LIGHT_BLUE, alignments=None,
              body_fills=None, font_size=9.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, header_fill)
        set_cell_border(
            cell,
            top={"val": "single", "sz": 6, "color": "B8C3CF"},
            bottom={"val": "single", "sz": 6, "color": "B8C3CF"},
            start={"val": "single", "sz": 6, "color": "B8C3CF"},
            end={"val": "single", "sz": 6, "color": "B8C3CF"},
        )
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(str(value))
        set_run_font(r, size=9.4, bold=True, color=NAVY)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        fill = body_fills[row_index] if body_fills and row_index < len(body_fills) else None
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if fill:
                set_cell_shading(cell, fill)
            set_cell_border(
                cell,
                top={"val": "single", "sz": 4, "color": "D5DBE2"},
                bottom={"val": "single", "sz": 4, "color": "D5DBE2"},
                start={"val": "single", "sz": 4, "color": "D5DBE2"},
                end={"val": "single", "sz": 4, "color": "D5DBE2"},
            )
            p = cell.paragraphs[0]
            alignment = (alignments[col_index] if alignments else
                         (WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 else WD_ALIGN_PARAGRAPH.LEFT))
            p.alignment = alignment
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)

    apply_table_geometry(
        table,
        widths,
        table_width_dxa=CONTENT_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_document(doc)
    bullet_num_id = add_numbering_definition(doc, ordered=False)

    add_title_block(doc)
    add_callout(
        doc,
        "一句话结论",
        "三个字段不是同一种判断：是否重点词看长期稳定性，优先级看短期机会与时间紧迫性，对应产品方向回答“具体可以做什么产品”。前两项已落为飞书表头公式；产品方向采用 AI 文本，最新版 Prompt 已完成，但尚无证据证明已在飞书保存并批量运行。",
        fill=PALE_GREEN,
        border="4F7C53",
    )

    doc.add_paragraph("1. 当前成果", style="Heading 1")
    add_table(
        doc,
        ["项目", "当前形式", "完成状态", "业务作用"],
        [
            ["是否重点词", "飞书公式 fₓ", "已迁移、已验证", "识别长期稳定出现且能形成产品需求的关键词"],
            ["优先级", "飞书公式 fₓ", "已迁移、已验证", "识别短期风口与跟进紧迫性，输出 A/B/C/观察中"],
            ["对应产品方向", "AI 文本", "Prompt 已完成；飞书保存与运行待确认", "把明确的关键词属性组合成可供运营参考的产品方向"],
            ["历史辅助字段", "数字字段", "3 个字段已新增", "为长期计数、近 8 批出现次数和批次顺序提供输入"],
        ],
        [1700, 1700, 2450, 3510],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                    WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        body_fills=[None, None, PALE_GOLD, None],
    )

    if SCREENSHOT.exists():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        picture = run.add_picture(str(SCREENSHOT), width=Inches(6.42))
        picture._inline.docPr.set(
            "descr",
            "飞书字段表头截图：是否重点词和优先级显示公式图标，对应产品方向显示文本字段图标",
        )
        picture._inline.docPr.set("title", "关键词经营决策三字段表头")
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(7)
        r = caption.add_run("图 1  当前表头：前两列的 fₓ 表示公式字段；对应产品方向为文本字段")
        set_run_font(r, size=9, italic=True, color=GRAY)

    add_label_paragraph(
        doc,
        "为什么网格里看不到公式文本",
        "飞书多维表格的公式配置在字段表头层级，不是逐个单元格填写。表头出现 fₓ 即表示整列由公式计算；新增记录也会自动套用。查看完整表达式需要打开字段配置。",
    )
    add_label_paragraph(
        doc,
        "当前结构快照",
        "主表 23 个字段、301 条记录，其中 300 条为有效关键词、1 条为空白记录；历史总表 19 个字段、300 条记录；关键词编号库 5 个字段、300 条记录。",
    )

    doc.add_paragraph("2. 三个字段如何分工", style="Heading 1")
    add_callout(
        doc,
        "核心原则",
        "长期性、紧迫性和产品语义必须拆开。否则会出现“新风口因为历史不够而错过”“稳定品牌词被误当成产品研发方向”“高热度改变了产品本身含义”等错误。",
        fill=CALLOUT,
        border=BLUE,
    )
    add_table(
        doc,
        ["判断轴", "回答的问题", "主要输入", "不承担的职责"],
        [
            ["是否重点词", "这个词是否长期、稳定地反复出现？", "关键词编号、有效批次数、近 8 批出现次数、分类与标签", "不判断短期风口；不因当期热度高就直接判“是”"],
            ["优先级", "现在要多快行动？", "分类、标签、用户意图、排名、搜索热度、交易热度、历史计数", "不等待重点词成熟；不定义产品长什么样"],
            ["对应产品方向", "关键词明确表达了什么产品方案？", "原始关键词、标准归并词、关键词分类、细分标签", "不决定是否值得做；不根据热度补造属性"],
        ],
        [1550, 2180, 3040, 2590],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_paragraph("联动方式", style="Heading 2")
    add_bullets(doc, [
        "是否重点词与优先级都读取关键词分类和细分标签，品牌词与清洁/漏水/排水/维修等服务型需求不会被误判为产品研发重点。",
        "优先级不依赖“是否重点词”的结果。新词即使还处于“观察中”，也可以因排名、搜索热度、交易热度和购买/对比意图被标记为 A，从而避免错过风口。",
        "对应产品方向只负责把明确属性组合成方向名称。运营使用时再用优先级筛选“先做什么”，用是否重点词判断“是否值得长期沉淀”。",
        "三个字段没有循环依赖，因此历史数据回填、热度变化或前置 AI 字段更新后，公式可以稳定重算。",
    ], bullet_num_id)

    doc.add_paragraph("运营阅读矩阵", style="Heading 2")
    add_table(
        doc,
        ["是否重点词", "优先级", "建议理解", "典型动作"],
        [
            ["是", "A / B", "长期核心需求且当前有机会", "优先进入产品规划或快速验证"],
            ["是", "C", "长期稳定，但当前没有明显窗口", "纳入常规产品池，持续跟踪"],
            ["观察中", "A", "历史不足，但出现明显风口信号", "立即小步测试，不必等满 8 批"],
            ["否", "A / B", "短期机会存在，但不属于长期产品核心", "限时测试，控制投入和库存"],
            ["任意", "任意", "产品方向为空", "不强行造产品；可能是品牌导航、店铺或服务问题"],
        ],
        [1650, 1450, 2960, 3300],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                    WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_paragraph("3. 是否重点词", style="Heading 1")
    add_label_paragraph(
        doc,
        "业务定义",
        "长期、稳定出现，并且能够形成产品需求的关键词。它不是“本周热度高”，也不是“所有经常出现的词”。品牌导航与典型售后服务词即使稳定出现，也不进入 V1 的产品重点词。",
    )
    add_label_paragraph(
        doc,
        "输出枚举",
        "空值、观察中、是、否。空搜索词输出空；历史不足时输出观察中。",
    )
    add_label_paragraph(
        doc,
        "长期门槛",
        "至少已有 8 个有效可比批次，并在最近 8 批中出现至少 6 次。按每周一次采集计算，确认“是”需要约 8 周，但这不会阻塞优先级提前发现风口。",
    )

    doc.add_paragraph("V1 判断顺序", style="Heading 2")
    add_numbered_steps(doc, [
        ("检查记录", "搜索词为空则返回空。"),
        ("检查历史成熟度", "已有有效批次数为空或少于 8，返回“观察中”。"),
        ("排除非产品型需求", "品牌词，或细分标签含痛点/清洁、痛点/漏水、痛点/排水、痛点/维修，返回“否”。"),
        ("计算稳定出现", "近 8 批出现次数大于等于 6 返回“是”，否则返回“否”。"),
    ])

    doc.add_paragraph("字段名版公式", style="Heading 2")
    add_code_block(doc, """
IF(搜索词为空, "",
  IF(已有有效批次数为空 OR 已有有效批次数<8, "观察中",
    IF(关键词分类="品牌词" OR 细分标签包含服务型痛点, "否",
      IF(近8批出现次数>=6, "是", "否")
    )
  )
)
""")
    add_callout(
        doc,
        "口径提醒",
        "“近 8 批出现 6 次”统计的是同一关键词编号在可比批次中的出现次数，同一批次重复记录只能算 1 次。关键词编号是跨周匹配的稳定键，不能用排名或当前行号代替。",
        fill=PALE_GOLD,
        border="A06B00",
    )

    doc.add_paragraph("4. 优先级", style="Heading 1")
    add_label_paragraph(
        doc,
        "业务定义",
        "优先级代表风口和时间紧迫性，回答“现在要不要快做”。它不是长期稳定性的替代，也不等于对应产品方向。",
    )
    add_table(
        doc,
        ["输出", "含义", "V1 触发逻辑"],
        [
            ["A-立即跟进", "短期窗口明确，需要快速验证", "至少 2 个有效批次；近 8 批仅出现 1 次；排名前 100；购买型或对比型；搜索热度高；交易热度中或高"],
            ["B-持续观察", "商业信号较强，但紧迫性尚不足", "购买型或对比型；搜索热度高；交易热度中或高"],
            ["C-常规跟踪", "没有短期高价值信号，或属于品牌/典型服务需求", "排除项直接为 C；其余未满足 A/B 的有效记录为 C"],
            ["观察中", "没有可用批次输入", "已有有效批次数为空或为 0"],
        ],
        [1600, 2400, 5360],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        body_fills=[PALE_RED, PALE_GOLD, None, LIGHT_GRAY],
    )

    doc.add_paragraph("为什么不等重点词确认后再给优先级", style="Heading 2")
    add_callout(
        doc,
        "防止错过风口",
        "重点词要看满 8 批才能确认长期性；优先级从第 1 个有效批次就可输出 B/C，从第 2 个有效批次开始可以识别 A。这样既保留长期判断的严谨性，也不要求运营等 8 周才行动。",
        fill=PALE_GREEN,
        border="4F7C53",
    )

    doc.add_paragraph("字段名版公式", style="Heading 2")
    add_code_block(doc, """
IF(搜索词为空, "",
  IF(已有有效批次数为空 OR 已有有效批次数=0, "观察中",
    IF(品牌词 OR 服务型痛点, "C-常规跟踪",
      IF(已有有效批次数>=2 AND 近8批出现次数=1 AND 排名<=100
         AND 用户意图为购买型/对比型 AND 搜索热度="高"
         AND 交易热度为"中"/"高", "A-立即跟进",
        IF(用户意图为购买型/对比型 AND 搜索热度="高"
           AND 交易热度为"中"/"高", "B-持续观察",
          "C-常规跟踪"
        )
      )
    )
  )
)
""")

    doc.add_paragraph("5. 对应产品方向", style="Heading 1")
    add_label_paragraph(
        doc,
        "业务定义",
        "把关键词中明确表达的产品属性组合成一个可参考的产品方向名称。允许留空，不为了填满字段而强行生成。",
    )
    add_label_paragraph(
        doc,
        "字段形式",
        "AI 文本，而不是公式。原因是产品方向涉及语言理解、属性保留和冲突判断，纯公式无法覆盖“高端”“深泡”“人造石”“小户型”等组合语义。",
    )
    add_label_paragraph(
        doc,
        "输入字段",
        "原始关键词、标准归并词、关键词分类、细分标签。原始关键词是最终事实依据；其余字段用于结构化理解，但不能覆盖原词。",
    )

    doc.add_paragraph("生成原则", style="Heading 2")
    add_bullets(doc, [
        "只保留原始关键词明确表达的属性，不使用常识补全。小户型不自动补深泡，老人不自动补防滑，高端不自动补人造石。",
        "不同产品属性不能抹平。高端浴缸与普通浴缸、深泡浴缸与通用浴缸是不同方向。",
        "方向名称按“定位 + 人群/场景 + 尺寸 + 材质 + 风格 + 形状/款式 + 安装方式 + 功能/需求 + 浴缸”组合，只保留有证据的部分。",
        "品牌、店铺、交易导航、地域和颜色不进入产品方向。",
        "明确痛点可转为产品需求，例如“浴缸太滑”可输出“防滑浴缸”；清洁、维修、漏水处理等服务问题留空。",
        "只有通用品类“浴缸”且没有其他产品属性时留空。",
    ], bullet_num_id)

    doc.add_paragraph("示例", style="Heading 2")
    add_table(
        doc,
        ["原始关键词", "AI 输出", "说明"],
        [
            ["家用浴缸", "家用浴缸", "保留明确场景"],
            ["高端人造石浴缸", "高端人造石浴缸", "定位和材质同时保留"],
            ["日式深泡小浴缸", "小型日式深泡浴缸", "保留尺寸、风格、功能"],
            ["浴缸太滑", "防滑浴缸", "明确痛点转为产品需求"],
            ["科勒浴缸官方旗舰店", "留空", "品牌/店铺导航不形成产品方向"],
            ["浴缸漏水维修", "留空", "服务问题不形成产品方向"],
        ],
        [2350, 2600, 4410],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_callout(
        doc,
        "当前未完成",
        "最新版产品方向 Prompt 已在本地定稿并通过规则测试；当前证据只能确认字段为文本类型，不能确认 Prompt 已成功保存到飞书，也没有触发批量 AI 运行。正式测试前需要在飞书 AI 字段配置中保存 Prompt，并先抽样运行。",
        fill=PALE_GOLD,
        border="A06B00",
    )

    doc.add_paragraph("6. 新增的三个辅助字段", style="Heading 1")
    add_callout(
        doc,
        "为什么只加三列",
        "V1 不增加搜索分、交易分、趋势分等大量隐藏列，只保留公式真正缺少的历史计数和批次顺序。业务视图可以把这些列隐藏，但自动化必须维护真实数值。",
        fill=CALLOUT,
        border=BLUE,
    )
    add_table(
        doc,
        ["字段", "所在表", "定义", "谁负责填写 / 用途"],
        [
            ["已有有效批次数", "关键词分析 V1（修正版）", "同一分析口径下已完成并可比较的批次数", "周更导入器回填；决定是否已有足够历史，以及优先级何时开始工作"],
            ["近8批出现次数", "关键词分析 V1（修正版）", "同一关键词编号在最近 8 个有效批次中出现的批次数；同批重复只算 1 次", "周更导入器回填；用于判断长期稳定与识别新出现/重新出现"],
            ["批次编号", "关键词历史总表 V1", "按同一口径递增的采集批次序号", "周更导入器写入；用于确定批次顺序，与采集日期共同保留"],
        ],
        [1650, 2050, 3100, 2560],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                    WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_label_paragraph(
        doc,
        "字段位置",
        "三个字段均追加在现有业务字段的最右侧：主表的两个辅助字段位于采集日期之后；历史总表的批次编号位于采集日期之后。运营日常视图可隐藏，避免干扰阅读。",
    )
    add_label_paragraph(
        doc,
        "有效批次口径",
        "每个批次只对应一个一级类目，并固定主关键词与来源渠道；采集日期和批次编号必须存在；同一周重复导入不应重复计数。可比较范围使用“一级类目 + 主关键词 + 来源渠道”，关键词跨批匹配使用永久关键词编号。",
    )

    doc.add_paragraph("当前为什么显示“观察中”", style="Heading 2")
    add_callout(
        doc,
        "不是公式失效",
        "历史表目前缺少可用的真实采集日期/批次输入，主表两个历史辅助字段也尚未由周更导入器回填。公式遇到空历史输入会按设计返回“观察中”，避免把未知数据误判成 A、B、是或否。",
        fill=PALE_GOLD,
        border="A06B00",
    )

    doc.add_paragraph("7. 每周自动联动流程", style="Heading 1")
    add_numbered_steps(doc, [
        ("生成本期快照", "一张新批次只对应一个一级类目；保留淘宝原始排名、区间和“-”。"),
        ("补齐批次信息", "写入真实采集日期和递增批次编号，不从文件名或当前时间猜测。"),
        ("绑定永久编号", "通过“一级类目 + 规范化原始关键词”查找关键词编号；同词同类目沿用编号。"),
        ("追加历史总表", "本期快照按原值追加，历史记录不覆盖。"),
        ("计算辅助字段", "按可比口径统计已有有效批次数，并按关键词编号统计近 8 批出现次数。"),
        ("回填主表", "只更新两个数字辅助字段；飞书自动重算是否重点词和优先级。"),
        ("运行产品方向 AI", "前置 AI 字段确认无误后抽样运行；结果可为空，不覆盖原始字段。"),
        ("运营验收", "优先检查 A 级词、重点词、产品方向为空及公式异常；确认后再扩大批量运行。"),
    ])

    doc.add_paragraph("自动更新边界", style="Heading 2")
    add_table(
        doc,
        ["变化", "是否重点词", "优先级", "对应产品方向"],
        [
            ["历史辅助数字更新", "自动重算", "自动重算", "不受影响"],
            ["排名/搜索热度/交易热度更新", "不受影响", "自动重算", "不受影响"],
            ["关键词分类/细分标签更新", "自动重算", "自动重算", "需要重新运行 AI 才会更新"],
            ["原始关键词/归并词更新", "通过相关前置字段间接影响", "通过相关前置字段间接影响", "需要重新运行 AI 才会更新"],
        ],
        [2500, 1950, 1950, 2960],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                    WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_callout(
        doc,
        "关键区别",
        "公式字段随输入变化自动重算；AI 文本字段不是普通公式，是否自动重新生成取决于飞书 AI 字段的触发设置。V1 建议显式抽样/批量运行，避免前置字段一改就消耗额度或覆盖已审核结果。",
        fill=CALLOUT,
        border=BLUE,
    )

    doc.add_paragraph("8. 验收结果与未完成项", style="Heading 1")
    add_table(
        doc,
        ["验收项", "结果", "说明"],
        [
            ["公式字段迁移", "通过", "是否重点词、优先级已由文本字段转换为飞书公式字段（type 20）"],
            ["隔离实表公式测试", "6 / 6 通过", "覆盖观察中、长期稳定、品牌排除、服务排除、新风口 A、普通 C"],
            ["本地自动化测试", "24 / 24 通过", "辅助字段、长期计数、优先级、产品方向 Prompt 和公式迁移验证"],
            ["其他字段与记录保护", "通过", "迁移前有完整备份；除两个目标字段类型/公式外，三张业务表与备份一致"],
            ["公式迁移写记录", "0 条", "通过字段级公式替换完成，没有逐行覆盖用户数据"],
            ["辅助字段创建", "完成", "主表新增 2 个数字字段，历史表新增 1 个数字字段"],
            ["历史计数自动回填", "未完成", "周更导入器尚未计算并写入已有有效批次数、近 8 批出现次数"],
            ["产品方向 AI 配置", "待确认", "Prompt 已完成；飞书保存、抽样生成和批量验收尚未完成"],
        ],
        [2380, 1700, 5280],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        body_fills=[PALE_GREEN, PALE_GREEN, PALE_GREEN, PALE_GREEN, PALE_GREEN,
                    PALE_GREEN, PALE_GOLD, PALE_GOLD],
    )

    doc.add_paragraph("正式跑测试前的最小动作", style="Heading 2")
    add_numbered_steps(doc, [
        ("完成周更回填", "让导入器按真实历史计算两个辅助字段；否则两个公式只能持续显示观察中。"),
        ("抽检公式结果", "使用实际关键词验证 A/B/C、是/否/观察中与运营理解一致。"),
        ("保存产品方向 Prompt", "在飞书 AI 字段中插入四个真实字段引用，先运行 10 至 20 条抽样。"),
        ("再做批量运行", "抽样无补造属性、无强行填充后，再对剩余记录运行。"),
    ])

    add_page_break(doc)
    doc.add_paragraph("附录 A：产品方向完整 Prompt", style="Heading 1")
    add_code_block(doc, """
你是浴缸产品方向分析助手。
原始关键词：{{原始关键词}}
标准归并词：{{标准归并词}}
关键词分类：{{关键词分类}}
细分标签：{{细分标签}}

目标：给运营一个可参考的产品方向名称。只输出方向名称，不解释；没有明确方向时留空。

判断依据：
1. 原始关键词是最终事实依据；标准归并词用于识别运营需求桶；关键词分类用于识别主属性；细分标签用于补齐原词中已经识别出的其他属性。
2. 字段冲突时以原始关键词为准。禁止补充原始关键词没有明确表达的属性，禁止根据常识推断。
3. 保留所有会改变产品方案的明确属性；同属性组合统一命名，不同属性不得合并。
4. 方向名称按“定位 + 人群/场景 + 尺寸 + 材质 + 风格 + 形状/款式 + 安装方式 + 功能/需求 + 浴缸”的顺序组合，只保留有证据的部分。
5. 品牌、店铺、交易导航、地域和颜色不进入产品方向。
6. 明确痛点可以直接转写为产品需求，例如“浴缸太滑”输出“防滑浴缸”。清洁、维修、漏水处理等服务问题不形成产品方向，留空。
7. 老人、户外、小户型等场景不得擅自补充防滑、耐候、深泡等属性；高端和普通、深泡浴缸和通用浴缸是不同方向。
8. 只有通用品类“浴缸”且没有其他产品属性时留空。

示例：
家用浴缸 -> 家用浴缸
高端人造石浴缸 -> 高端人造石浴缸
日式深泡小浴缸 -> 小型日式深泡浴缸
科勒浴缸官方旗舰店 -> 留空
浴缸漏水维修 -> 留空

只输出一个方向名称或留空。
""")

    doc.add_paragraph("附录 B：公式验收样例", style="Heading 1")
    add_table(
        doc,
        ["场景", "关键输入摘要", "是否重点词", "优先级"],
        [
            ["历史不足", "7 批 / 出现 7 次；高搜索、中交易", "观察中", "B-持续观察"],
            ["长期稳定", "8 批 / 出现 6 次；购买型；高搜索、中交易", "是", "B-持续观察"],
            ["品牌导航", "品牌词；8 批 / 出现 8 次", "否", "C-常规跟踪"],
            ["服务需求", "痛点/清洁；8 批 / 出现 8 次", "否", "C-常规跟踪"],
            ["新风口", "2 批 / 出现 1 次；排名 99；对比型；高搜索、中交易", "观察中", "A-立即跟进"],
            ["普通低信号", "8 批 / 出现 2 次；低搜索、低交易", "否", "C-常规跟踪"],
        ],
        [1550, 4210, 1700, 1900],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                    WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=9.1,
    )
    add_callout(
        doc,
        "文档边界",
        "本稿说明的是 2026-08-13 已落地的 V1 公式与辅助字段合同。阈值和排除口径可在真实周更数据跑出后由运营复审，但调整必须同时更新公式、测试样例和说明，不能只改表内结果值。",
        fill=LIGHT_GRAY,
        border="718096",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
