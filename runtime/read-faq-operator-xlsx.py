import argparse
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook


def values(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ValueError(f"Worksheet {sheet.title} is empty")
    headers = [str(value).strip() if value is not None else "" for value in headers]
    if any(not header for header in headers):
        raise ValueError(f"Worksheet {sheet.title} has an empty header")
    return headers, [dict(zip(headers, row)) for row in rows]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    args = parser.parse_args()
    workbook = load_workbook(args.xlsx, data_only=True, read_only=True)
    required = ["痛点深度分析"]
    missing = [name for name in required if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Missing worksheets: {', '.join(missing)}")
    headers, rows = values(workbook["痛点深度分析"])
    expected = ["痛点类型", "出现次数", "占比", "痛点描述", "典型用户原话"]
    if headers != expected:
        raise ValueError(f"Worksheet 痛点深度分析 headers do not match: {headers}")
    print(json.dumps({"sheet": "痛点深度分析", "headers": headers, "rows": rows}, ensure_ascii=False))


if __name__ == "__main__":
    main()
