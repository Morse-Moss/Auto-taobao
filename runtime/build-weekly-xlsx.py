#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把自采的周报 CSV 转成带内嵌图的 XLSX。

为什么必须带内嵌图：飞书的「商品图片」是附件字段（type 17），导入器
skills/xws-to-feishu-base/scripts/extract_xws_xlsx.py 是靠 openpyxl 读
sheet._images + anchor._from 把图抽出来再上传换 file_token 的。
CSV 里的图片 URL 用不上——它不会去下载。

格式照 2026-10-19 小旺神那份 xlsx 复刻（实测参数）：
  sheet 名   市场分析
  图片       220×220 像素，锚在 B 列（anchor.col == 1），JPEG
  列宽行高   都不设（那份也没有设，少一个变量）
  行数       表头 1 行 + 数据 N 行

与 extract_xws_xlsx.py 的契约对齐点（改这里之前先读那个脚本）：
  - 必须是一个 sheet 且在 active
  - 表头第 1 行，每列都要是非空字符串
  - 图片必须真有 anchor（用 ws.add_image(img, 'B2') 就有，别手动写 drawing）

用法：
  python runtime/build-weekly-xlsx.py --csv a.csv --images <dir> --out a.xlsx [--manifest m.json]

图片目录约定：文件名是 <item_id>.<ext>，item_id 从 CSV 的商品链接里取。
缺图的商品不会中断生成，但会记进 manifest，由调用方决定是否接受。
"""
import argparse
import csv
import io
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

SHEET_TITLE = "市场分析"
THUMB_SIZE = (220, 220)          # 与 10-19 那份一致
IMAGE_COLUMN = "B"                # 商品图片是第 2 列
EXPECTED_COLUMNS = 16


def die(message: str) -> None:
    print(json.dumps({"ok": False, "error": message}, ensure_ascii=False))
    sys.exit(1)


def item_id_of(link: str) -> str:
    """从 https://item.taobao.com/item.htm?id=123 里取 123。取不到返回空串。"""
    marker = "id="
    index = link.find(marker)
    if index < 0:
        return ""
    tail = link[index + len(marker):]
    digits = ""
    for ch in tail:
        if ch.isdigit():
            digits += ch
        else:
            break
    return digits


def to_thumbnail(raw: bytes):
    """把任意格式的图转成 220×220 的缩略图。

    返回 (PIL.Image, 'jpeg' | 'png')。
    透明图（PNG 带 alpha）垫白底再转 JPEG——淘宝主图基本是白底，
    直接转会把透明区变黑，看着像坏图。
    """
    source = PILImage.open(io.BytesIO(raw))
    source.load()

    has_alpha = source.mode in ("RGBA", "LA") or (
        source.mode == "P" and "transparency" in source.info
    )
    if has_alpha:
        source = source.convert("RGBA")
        canvas = PILImage.new("RGB", source.size, (255, 255, 255))
        canvas.paste(source, mask=source.split()[-1])
        flat = canvas
        target_format = "jpeg"
    else:
        flat = source.convert("RGB")
        # 原生 PNG 保留 PNG，避免二次压缩损失；其余一律 JPEG
        target_format = "png" if source.format == "PNG" else "jpeg"

    thumbnail = flat.resize(THUMB_SIZE, PILImage.LANCZOS)
    return thumbnail, target_format


def build(csv_path: Path, image_dir: Path, out_path: Path, manifest_path: Path | None) -> dict:
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            headers = next(reader)
        except StopIteration:
            die(f"CSV 是空的：{csv_path}")
        rows = [row for row in reader if any(cell.strip() for cell in row)]

    if len(headers) != EXPECTED_COLUMNS:
        die(f"表头应为 {EXPECTED_COLUMNS} 列，实际 {len(headers)} 列：{headers}")
    if any(not isinstance(h, str) or not h.strip() for h in headers):
        die(f"表头存在空列名：{headers}")

    # 注意：index() 给的是 0 基下标，而下面写单元格用的 column 是 1 基。
    # 这里必须 +1，否则「跳过图片列」会跳错一列——实测代价是序号整列丢空、
    # 图片列残留 URL，下游 validateDataset 直接判死（序号必须连续 1..N）。
    image_col_index = headers.index("商品图片") + 1 if "商品图片" in headers else None
    link_col_index = headers.index("商品链接") if "商品链接" in headers else None
    if image_col_index is None or link_col_index is None:
        die("CSV 缺少「商品图片」或「商品链接」列")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    for column, header in enumerate(headers, 1):
        sheet.cell(row=1, column=column, value=header)

    missing = []
    embedded = 0
    failed = []

    for row_index, row in enumerate(rows, start=2):
        # 补齐短行，保证每行都是 16 列——openpyxl 不会因为少一格而报错，
        # 但下游按列名取值时会静默拿到 None，宁可在这里显式补空。
        padded = list(row) + [""] * (len(headers) - len(row))
        for column, value in enumerate(padded, 1):
            # 图片列不写 URL：飞书那边要的是附件，写 URL 只会变成一列没用的文本。
            if column == image_col_index:
                sheet.cell(row=row_index, column=column, value=None)
                continue
            sheet.cell(row=row_index, column=column, value=value if value != "" else None)

        item_id = item_id_of(padded[link_col_index])
        if not item_id:
            missing.append({"row": row_index, "itemId": "", "reason": "no_item_id"})
            continue

        candidates = list(image_dir.glob(f"{item_id}.*"))
        if not candidates:
            missing.append({"row": row_index, "itemId": item_id, "reason": "file_absent"})
            continue

        raw = candidates[0].read_bytes()
        try:
            thumbnail, fmt = to_thumbnail(raw)
        except Exception as error:  # 损坏图不该让整批失败
            failed.append({"row": row_index, "itemId": item_id, "error": str(error)[:200]})
            continue

        buffer = io.BytesIO()
        if fmt == "png":
            thumbnail.save(buffer, format="PNG", optimize=True)
        else:
            thumbnail.save(buffer, format="JPEG", quality=85, optimize=True)
        buffer.seek(0)

        picture = XLImage(buffer)
        picture.width = THUMB_SIZE[0]
        picture.height = THUMB_SIZE[1]
        # 锚点必须用 add_image 建立，别手写 drawing xml——extract 那边靠 anchor 定位行号
        sheet.add_image(picture, f"{IMAGE_COLUMN}{row_index}")
        embedded += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)

    # 生成后自检：列错位这类 bug 在 build 里看不出来，要等下游 extract 才暴露。
    # 两个不变量：图片列不能有文本（飞书要附件）、序号必须连续 1..N。
    if "商品图片" in headers:
        leaked = sum(
            1 for row_index in range(2, len(rows) + 2)
            if str(sheet.cell(row=row_index, column=image_col_index).value or "").strip()
        )
        if leaked:
            die(f"自检失败：{leaked} 行的「商品图片」单元格残留文本，应为空")
    if "序号" in headers:
        serial_col = headers.index("序号") + 1
        for offset, row_index in enumerate(range(2, len(rows) + 2)):
            actual = str(sheet.cell(row=row_index, column=serial_col).value or "").strip()
            if actual != str(offset + 1):
                die(f"自检失败：第 {row_index} 行序号为 {actual!r}，应为 {offset + 1}")

    workbook.save(out_path)

    result = {
        "ok": True,
        "xlsx": str(out_path),
        "sheet": SHEET_TITLE,
        "rows": len(rows),
        "columns": len(headers),
        "embeddedImages": embedded,
        "missingImages": len(missing) + len(failed),
        "missing": missing[:200],
        "failed": failed[:50],
        "bytes": out_path.stat().st_size,
    }

    if manifest_path:
        manifest_path.parent.mkdir(parents=True, exist_ok=True)
        manifest_path.write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv", type=Path, required=True)
    parser.add_argument("--images", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, default=None)
    args = parser.parse_args()

    if not args.csv.exists():
        die(f"CSV 不存在：{args.csv}")
    if not args.images.is_dir():
        die(f"图片目录不存在：{args.images}")

    result = build(args.csv, args.images, args.out, args.manifest)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
